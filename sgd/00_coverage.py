"""
00_coverage.py -- the sgd piece's coverage listing for raw/. Prints labels and
coverage only, never an observation value.

  python3 sgd/00_coverage.py [--raw DIR]

For every file in raw/ (except RETRIEVED.txt): bytes and md5. Then, by type:

  .csv   long format (a period column such as TIME_PERIOD): the column
         headers; one line per series (the text of its identifying columns:
         codes and labels), with the number of periods, the first and last
         period, and the count of empty observations. Wide format (period
         headers, one row per series, as data.gov.sg serves): the same, per
         row. Columns that are mostly numeric are never printed: they are
         values, not labels.
         Observation-level attribute columns (OBS_STATUS, OBS_CONF, ...) are
         not part of a series' identity; their distinct codes are listed
         once.
  .json  SingStat TableBuilder: title, dataLastUpdated, and per row its
         seriesNo, rowText and unit, with the period count, first and last
         period key. MAS chart API ({"name", "elements": [...]}): the name,
         the element keys, the count, the first and last date, and the
         distinct "updatedat" stamps.
  .zip   the member names and sizes; CSV members as above.
  .xlsx  sheet names and their row and column counts (openpyxl, read-only).
  .html  the page <title> only.
  .pdf   the /Title metadata only.

Standard library only, apart from openpyxl for .xlsx. Writes nothing.
"""
import argparse
import csv
import hashlib
import io
import json
import os
import re
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
PERIOD_NAMES = {"time_period", "period", "date", "month", "quarter", "year", "time"}
VALUE_NAMES = {"obs_value", "value", "values"}


def md5(data):
    return hashlib.md5(data).hexdigest()


def ascii_text(s):
    """ASCII only, whitespace collapsed (non-ASCII characters dropped)."""
    return " ".join(s.encode("ascii", "ignore").decode("ascii").split())


def period_key(s):
    """A sortable key for the period strings the sources use, or None."""
    t = s.strip()
    m = re.fullmatch(r"(\d{4})", t)
    if m:
        return (int(m.group(1)), 0, 0)
    m = re.fullmatch(r"(\d{4})[- ]?(?:M)?(\d{1,2})", t)
    if m and 1 <= int(m.group(2)) <= 12:
        return (int(m.group(1)), int(m.group(2)), 0)
    m = re.fullmatch(r"(\d{4})[- ]?([A-Za-z]{3})[a-z]*", t)
    if m and m.group(2).lower() in MONTHS:
        return (int(m.group(1)), MONTHS[m.group(2).lower()], 0)
    m = re.fullmatch(r"(\d{4})[- ]?(?:Q([1-4])|([1-4])Q)", t)
    if m:
        return (int(m.group(1)), 3 * int(m.group(2) or m.group(3)), 0)
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", t)
    if m:
        return (int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.fullmatch(r"(\d{4})-W(\d{2})", t)
    if m:
        return (int(m.group(1)), 0, int(m.group(2)))
    return None


def span(periods):
    """(count, first, last) of distinct periods, in time order if parseable."""
    ps = sorted(set(p for p in periods if p.strip()))
    keyed = [(period_key(p), p) for p in ps]
    if keyed and all(k is not None for k, _ in keyed):
        keyed.sort()
        return len(ps), keyed[0][1], keyed[-1][1], "time order"
    return len(ps), (ps[0] if ps else ""), (ps[-1] if ps else ""), "string order (unparsed)"


def numeric(s):
    try:
        float(s.replace(",", ""))
        return True
    except ValueError:
        return False


def csv_listing(text, indent="  "):
    rows = list(csv.reader(io.StringIO(text)))
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        print(indent + "empty")
        return
    head, body = rows[0], rows[1:]
    print(indent + f"rows {len(body)}, columns {len(head)}")
    pcols = [i for i, h in enumerate(head) if h.strip().lower() in PERIOD_NAMES]
    wide = [i for i, h in enumerate(head) if period_key(h) is not None]
    if not pcols and len(wide) >= 3:
        print(indent + "layout  wide (period headers)")
        lab = [i for i in range(len(head)) if i not in wide]
        lab = [i for i in lab
               if sum(numeric(r[i]) for r in body if i < len(r) and r[i].strip())
               <= 0.5 * max(1, sum(1 for r in body if i < len(r) and r[i].strip()))]
        print(indent + "label columns  " + " | ".join(head[i] for i in lab))
        for r in body:
            filled = [head[i] for i in wide if i < len(r) and r[i].strip() not in ("", "na", "NA", "-")]
            n, a, b, how = span(filled)
            print(indent + "  " + " | ".join(r[i] for i in lab if i < len(r))
                  + f"  :: {n} periods, {a} to {b} ({how}), {len(wide) - n} empty")
        return
    if not pcols:
        print(indent + "layout  no period column found; headers only")
        print(indent + "columns  " + " | ".join(head))
        return
    p = pcols[0]
    vcols = [i for i, h in enumerate(head) if h.strip().lower() in VALUE_NAMES]
    mostly_numeric = set()
    for i in range(len(head)):
        vals = [r[i] for r in body if i < len(r) and r[i].strip()]
        if vals and sum(numeric(v) for v in vals) > 0.5 * len(vals):
            mostly_numeric.add(i)
    attrs = [i for i, h in enumerate(head) if h.strip().upper().startswith("OBS_") and i not in vcols]
    lab = [i for i in range(len(head))
           if i != p and i not in vcols and i not in mostly_numeric and i not in attrs]
    print(indent + "layout  long (period column " + head[p] + ")")
    print(indent + "columns  " + " | ".join(head))
    print(indent + "not printed (values or numeric)  "
          + (" | ".join(head[i] for i in sorted((set(vcols) | mostly_numeric) - {p})) or "none"))
    for i in attrs:
        if i in mostly_numeric:
            continue
        codes = sorted({r[i] for r in body if i < len(r)} - {""})
        print(indent + f"attribute {head[i]} codes  " + (", ".join(codes) if codes else "(always empty)"))
    series = {}
    for r in body:
        k = tuple(r[i] if i < len(r) else "" for i in lab)
        empty = any(not (r[v].strip() if v < len(r) else "") or r[v].strip().lower() in ("nan", "na")
                    for v in vcols)
        series.setdefault(k, []).append((r[p] if p < len(r) else "", empty))
    varying = [j for j in range(len(lab)) if len({k[j] for k in series}) > 1]
    const = [j for j in range(len(lab)) if j not in varying]
    if const:
        k0 = next(iter(series))
        print(indent + "same for every series  " + " | ".join(f"{head[lab[j]]}={k0[j]}" for j in const))
    print(indent + f"series {len(series)}")
    for k, obs in sorted(series.items()):
        n, a, b, how = span([o for o, _ in obs])
        print(indent + "  " + " | ".join(k[j] for j in varying)
              + f"  :: {n} periods, {a} to {b} ({how}), {sum(e for _, e in obs)} empty")


def json_listing(data, indent="  "):
    try:
        d = json.loads(data.decode("utf-8-sig"))
    except ValueError as e:
        print(indent + f"not JSON: {e}")
        return
    if isinstance(d, dict) and isinstance(d.get("elements"), list):
        els = [e for e in d["elements"] if isinstance(e, dict)]
        keys = sorted({k for e in els for k in e})
        n, a, b, how = span([str(e.get("date", "")) for e in els])
        print(indent + f"layout  MAS chart API, name {d.get('name')!r}")
        print(indent + "element keys  " + ", ".join(keys) + "  (value not printed)")
        print(indent + f"elements {len(els)}  :: {n} distinct dates, {a} to {b} ({how}), "
              f"{sum(e.get('value') in (None, '') for e in els)} empty")
        stamps = sorted({str(e.get("updatedat")) for e in els})
        print(indent + f"updatedat  {len(stamps)} distinct: " + ", ".join(stamps[:3])
              + (" ..." if len(stamps) > 3 else ""))
        return
    D = d.get("Data") if isinstance(d, dict) else None
    if not isinstance(D, dict) or "row" not in D:
        print(indent + "JSON, not the TableBuilder layout; top-level keys: "
              + (", ".join(d.keys()) if isinstance(d, dict) else type(d).__name__))
        return
    print(indent + f"title   {D.get('title')}")
    print(indent + f"id      {D.get('id')}")
    print(indent + f"last updated {D.get('dataLastUpdated')}")
    print(indent + f"series  {len(D['row'])}")
    for r in D["row"]:
        cols = r.get("columns", [])
        n, a, b, how = span([c.get("key", "") for c in cols])
        print(indent + f"  {r.get('seriesNo')} | {r.get('rowText')} | {r.get('uoM')}"
              + f"  :: {n} periods, {a} to {b} ({how})")


def xlsx_listing(path, indent="  "):
    try:
        import openpyxl
    except ImportError:
        print(indent + "openpyxl not installed: sheet listing skipped")
        return
    wb = openpyxl.load_workbook(path, read_only=True)
    for ws in wb.worksheets:
        print(indent + f"sheet {ascii_text(ws.title)!r}: {ws.max_row} rows, {ws.max_column} columns")


def main():
    ap = argparse.ArgumentParser(description="sgd step 00: coverage listing of raw/ (no values)")
    ap.add_argument("--raw", default=os.path.join(HERE, "raw"))
    a = ap.parse_args()
    files = sorted(f for f in os.listdir(a.raw)
                   if f != "RETRIEVED.txt" and not f.startswith(".") and os.path.isfile(os.path.join(a.raw, f)))
    if not files:
        print("raw/ holds no data file yet (only RETRIEVED.txt).")
        return
    for f in files:
        path = os.path.join(a.raw, f)
        data = open(path, "rb").read()
        print(f"{f}\n  bytes   {len(data)}\n  md5     {md5(data)}")
        ext = f.lower().rsplit(".", 1)[-1]
        if ext == "csv":
            csv_listing(data.decode("utf-8-sig", errors="replace"))
        elif ext == "json":
            json_listing(data)
        elif ext == "zip":
            with zipfile.ZipFile(path) as z:
                for m in z.infolist():
                    print(f"  member {m.filename} ({m.file_size} bytes)")
                    if m.filename.lower().endswith(".csv"):
                        csv_listing(z.read(m).decode("utf-8-sig", errors="replace"), indent="    ")
        elif ext == "xlsx":
            xlsx_listing(path)
        elif ext in ("html", "htm"):
            m = re.search(rb"<title[^>]*>\s*(.*?)\s*</title>", data, re.S | re.I)
            print("  title   " + (ascii_text(m.group(1).decode("utf-8", "replace")) if m else "(none)"))
        elif ext == "pdf":
            m = re.search(rb"/Title\s*\((.*?)\)", data)
            raw = m.group(1).decode("latin-1") if m else ""
            ok = raw and sum(32 <= ord(c) < 127 for c in raw) >= 0.9 * len(raw)
            print("  title   " + (ascii_text(raw) if ok else
                                  "(/Title metadata not plain text)" if m else "(no /Title metadata)"))
        print()


if __name__ == "__main__":
    sys.exit(main())
