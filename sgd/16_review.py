"""
16_review.py -- analyses added after outside review, NOT SEALED and NOT
SCORED (THESIS_ADDENDUM item 10). Changes no sealed number, outcome,
threshold or verdict. Reuses the sealed definitions in 11_tests.py (the log
cross rate, the three-month endpoint means, the split, Spearman with average
ranks, the interval table) so that every number here is on the same footing.

  python3 sgd/16_review.py [--root DIR]

Reads out/ (10_load.py, 11_tests.py), raw/s1c (the BIS broad-basket weights)
and raw/s8 (Japan's CPI, BIS long CPI series). Writes, each with columns
key, value, printed, meaning (14_manifest.py puts every printed value into
the number manifest):

  out/review_a1_usd.csv      A1. US-dollar numeraire split, an exact identity:
                             b = dln(X per USD) - dln(SGD per USD).
  out/review_a2_expair.csv   A2. Ex-pair split: each currency removed from
                             the other's broad basket, a first-order
                             approximation for a chain-linked index.
  out/review_a3_real.csv     A3. Real split of the yen cross:
                             b_real = b + dln CPI_SG - dln CPI_JP, against the
                             BIS real broad indices.
  out/review_a4_rolling.csv  A4. Rolling five-year windows for the yen: the
                             counts.
  out/review_a5_t7.csv       A5. T7: a moving-block bootstrap for D; growth
                             and CPI inflation led by 1 and 2 quarters; the
                             effective sample size.
  out/review_a6_context.csv  Context for the article (round 5): USD/JPY at the
                             scored window's endpoints (S2), and where the
                             BIS real broad index for Japan (S1b) sat in its
                             series since 1994.
  out/review_article.csv     Article forms of numbers already sealed or above.

and out/review_rolling_windows.csv, one row per rolling window (the data for
chart 4; not printed).
"""
import csv
import importlib.util
import io
import math
import os
import random

import openpyxl

import sgdlib as L

_spec = importlib.util.spec_from_file_location("t11", os.path.join(L.SGD, "11_tests.py"))
T11 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(T11)

F_WEIGHTS = "s1c_bis_eer_weights_broad.xlsx"
F_CPI_JP = "s8_bis_long_cpi_jp_monthly.csv"
BLOCKS = (2, 4, 8)          # moving-block lengths, in intervals; 4 is the main one
LEADS = (1, 2)              # quarters
ROLL_MONTHS, ROLL_STEP = 60, 3


class Rows:
    def __init__(self):
        self.rows = []

    def add(self, key, value, fmt, meaning):
        if value is None:
            return
        self.rows.append([key, value, fmt.format(value), meaning])

    def pct(self, key, x, meaning):
        """Per-cent move of log change x, one decimal and whole, absolute."""
        if x is None:
            return
        v = 100 * (math.exp(x) - 1)
        word = "up" if v >= 0 else "down"
        self.add(key + "_pct", abs(v), "{:.1f}", f"{meaning}, per cent, {word}")
        self.add(key + "_pct0", abs(v), "{:.0f}", f"{meaning}, per cent, {word}, whole")

    def share(self, key, x, meaning):
        if x is None:
            return
        self.add(key, x, "{:.4f}", meaning + ", share of the log change")
        self.add(key + "_pct0", abs(100 * x), "{:.0f}",
                 meaning + ", per cent of the log change, whole" + ("" if x >= 0 else ", negative"))

    def write(self, path):
        buf = io.StringIO(newline="")
        w = csv.writer(buf, lineterminator="\n")
        w.writerow(["key", "value", "printed", "meaning"])
        for key, value, printed, meaning in self.rows:
            w.writerow([key, L.fmt(value), printed, meaning])
        L.write_text(path, buf.getvalue())
        print(f"  {os.path.relpath(path, os.path.dirname(os.path.dirname(path)))}: {len(self.rows)} rows")


def change(f, start, end):
    a, z = T11.endmean(f, start), T11.endmean(f, end)
    return None if a is None or z is None else z - a


# ------------------------------------------------------------------ A1
def a1(D, W):
    R = Rows()
    for w, (first, last, start, end) in W.items():
        for cur in L.QUESTION:
            dsg = change(lambda m: math.log(D["usd"]["SGD"][m]) if m in D["usd"]["SGD"] else None, start, end)
            dx = change(lambda m: math.log(D["usd"][cur][m]) if m in D["usd"][cur] else None, start, end)
            if dsg is None or dx is None:
                continue
            b = dx - dsg
            k = f"A1_{w}_{cur}"
            R.add(k + "_b", b, "{:.4f}", f"{w} window: ln(X per SGD) change, = d ln(X per USD) - d ln(SGD per USD)")
            R.pct(k + "_sgd_vs_usd", -dsg, f"{w} window: the Singapore dollar against the US dollar")
            R.pct(k + "_x_vs_usd", -dx, f"{w} window: {cur} against the US dollar")
            if b > 0:
                R.share(k + "_share_sgd", -dsg / b, f"{w} window, {cur}: the Singapore dollar's own move against the US dollar")
                R.share(k + "_share_x", dx / b, f"{w} window, {cur}: {cur}'s own fall against the US dollar")
    return R


# ------------------------------------------------------------------ A2
def weight_sheets(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    out = {}
    for sn in wb.sheetnames:
        y0, y1 = (int(x) for x in sn.split("_"))
        rows = list(wb[sn].iter_rows(values_only=True))
        hdr = rows[5]
        mat = {}
        for r in rows[6:]:
            if r[1]:
                mat[r[1]] = {hdr[j]: r[j] for j in range(2, len(hdr)) if isinstance(r[j], (int, float))
                             and hdr[j] != "Total"}
        out[(y0, y1)] = mat
    return out


def window_weight(sheets, months_, area, partner):
    """Weight of `partner` in `area`'s basket, as a share, averaged over the
    window's months; each month takes the sheet whose three-year period holds
    its year, and years after the last sheet take the last sheet."""
    keys = sorted(sheets)
    ws = []
    for m in months_:
        y = int(m[:4])
        k = next((k for k in keys if k[0] <= y <= k[1]), keys[-1] if y > keys[-1][1] else keys[0])
        ws.append(sheets[k][area].get(partner, 0.0) / 100)
    return sum(ws) / len(ws)


def a2(D, W, raw):
    R = Rows()
    path = os.path.join(raw, F_WEIGHTS)
    if not os.path.exists(path):
        return R
    sheets = weight_sheets(path)
    for w, (first, last, start, end) in W.items():
        ms = L.months(first, last)
        for cur in L.QUESTION:
            area = L.CUR_AREA[cur]
            sp = T11.split(D, cur, start, end)
            if not sp or sp["b"] <= 0:
                continue
            b, s, nx = sp["b"], sp["s"], sp["nx"]
            w_sg = window_weight(sheets, ms, "SG", area)     # X in Singapore's basket
            w_x = window_weight(sheets, ms, area, "SG")      # SGD in X's basket
            # SG's index against X moves by b; X's index against SGD moves by -b.
            s_ex = (s - w_sg * b) / (1 - w_sg)
            n_ex = (nx + w_x * b) / (1 - w_x)
            k = f"A2_{w}_{cur}"
            R.add(k + "_w_in_sg", 100 * w_sg, "{:.1f}", f"{w} window: weight of {cur} in Singapore's basket, per cent")
            R.add(k + "_w_sg_in_x", 100 * w_x, "{:.1f}", f"{w} window: weight of SGD in {area}'s basket, per cent")
            R.add(k + "_s_ex", s_ex, "{:.4f}", f"{w} window: SGD broad index without {cur}, log change")
            R.add(k + "_n_ex", n_ex, "{:.4f}", f"{w} window: {area} broad index without SGD, log change")
            R.share(k + "_S", s_ex / b, f"{w} window, {cur}, ex-pair: the Singapore dollar rising")
            R.share(k + "_P", -n_ex / b, f"{w} window, {cur}, ex-pair: {cur} falling")
            R.share(k + "_R", (b - s_ex + n_ex) / b, f"{w} window, {cur}, ex-pair: the remainder")
    return R


# ------------------------------------------------------------------ A3
def load_cpi_jp(raw):
    path = os.path.join(raw, F_CPI_JP)
    if not os.path.exists(path):
        return {}
    out = {}
    for r in L.read_csv(path):
        if r["REF_AREA"] == "JP" and r["UNIT_MEASURE"] == "628" and r["OBS_VALUE"]:
            out[r["TIME_PERIOD"]] = float(r["OBS_VALUE"])
    return out


def a3(D, W, cpi_jp):
    R = Rows()
    if not cpi_jp:
        return R
    lncpi = lambda d: (lambda m: math.log(d[m]) if m in d else None)
    for w, (first, last, start, end) in W.items():
        sp = T11.split(D, "JPY", start, end)
        sg, jp = change(lncpi(D["cpi"]), start, end), change(lncpi(cpi_jp), start, end)
        s_r = change(lambda m: T11.ln_eer(D, "SG", m, "reer"), start, end)
        n_r = change(lambda m: T11.ln_eer(D, "JP", m, "reer"), start, end)
        if None in (sp, sg, jp, s_r, n_r):
            continue
        br = sp["b"] + sg - jp
        k = f"A3_{w}_JPY"
        R.pct(k + "_cpi_sg", sg, f"{w} window: Singapore's CPI (S7)")
        R.pct(k + "_cpi_jp", jp, f"{w} window: Japan's CPI (BIS long CPI)")
        R.add(k + "_b_real", br, "{:.4f}", f"{w} window: real yen per SGD, log change")
        R.pct(k + "_b_real", br, f"{w} window: the Singapore dollar against the yen, real")
        R.pct(k + "_s_real", s_r, f"{w} window: the Singapore dollar's real broad index")
        R.pct(k + "_n_real", n_r, f"{w} window: the yen's real broad index")
        # A traveller's figure (round 5): Tokyo goods per Singapore dollar,
        # the nominal cross deflated by Japan's prices alone.
        R.pct(k + "_tokyo_goods", sp["b"] - jp, f"{w} window: Tokyo goods per Singapore dollar (b - dln CPI_JP)")
        if br > 0:
            R.share(k + "_S", s_r / br, f"{w} window, real: the Singapore dollar rising")
            R.share(k + "_P", -n_r / br, f"{w} window, real: the yen falling")
            R.share(k + "_R", (br - s_r + n_r) / br, f"{w} window, real: the remainder")
    return R


# ------------------------------------------------------------------ A4
def a4(D, out):
    last = T11.last_month(D)
    rows, m = [], L.LONG_FIRST
    while L.month_add(m, ROLL_MONTHS - 1) <= last:
        z = L.month_add(m, ROLL_MONTHS - 1)
        start = tuple(L.month_add(m, k) for k in range(3))
        end = tuple(L.month_add(z, k) for k in (-2, -1, 0))
        sp = T11.split(D, "JPY", start, end)
        if sp:
            rows.append({"start": m, "end": z, "b": sp["b"], "s": sp["s"], "nx": sp["nx"],
                         "S": sp["S"], "P": sp["P"], "rose": int(sp["b"] > 0),
                         "yen_larger": int(sp["b"] > 0 and sp["P"] > sp["S"])})
        m = L.month_add(m, ROLL_STEP)
    buf = io.StringIO(newline="")
    w = csv.writer(buf, lineterminator="\n")
    cols = ["start", "end", "b", "s", "nx", "S", "P", "rose", "yen_larger"]
    w.writerow(cols)
    for r in rows:
        w.writerow([r[c] if c in ("start", "end", "rose", "yen_larger") else L.fmt(r[c]) for c in cols])
    L.write_text(os.path.join(out, "review_rolling_windows.csv"), buf.getvalue())
    R = Rows()
    rose = [r for r in rows if r["rose"]]
    R.add("A4_windows", len(rows), "{:d}", "rolling five-year windows, August 2005 on, in three-month steps")
    R.add("A4_rose", len(rose), "{:d}", "windows in which the Singapore dollar rose against the yen")
    R.add("A4_not_rose", len(rows) - len(rose), "{:d}", "windows in which it did not")
    R.add("A4_yen_larger", sum(r["yen_larger"] for r in rose), "{:d}",
          "of the windows where it rose, those where the yen's own fall was the larger part")
    R.add("A4_sgd_larger", sum(1 for r in rose if not r["yen_larger"]), "{:d}",
          "of the windows where it rose, those where the Singapore dollar's own rise was larger")
    if rows:
        R.add("A4_first_start_year", int(rows[0]["start"][:4]), "{:d}", "first window starts (year)")
        R.add("A4_last_end_year", int(rows[-1]["end"][:4]), "{:d}", "last window ends (year)")
    return R


# ------------------------------------------------------------------ A5
def block_boot(iv, block, seed=L.BOOT_SEED, draws=L.BOOT_N):
    rng = random.Random(seed)
    n, ds = len(iv), []
    nb = math.ceil(n / block)
    for _ in range(draws):
        smp = []
        for _ in range(nb):
            k = rng.randrange(n - block + 1)
            smp.extend(iv[k:k + block])
        _, _, d, _ = T11.rho_d(smp[:n])
        if d is not None:
            ds.append(d)
    ds.sort()

    def q(p):
        k = p * (len(ds) - 1)
        lo = int(math.floor(k))
        hi = min(lo + 1, len(ds) - 1)
        return ds[lo] + (ds[hi] - ds[lo]) * (k - lo)
    return q(0.05), q(0.95)


def lag1(xs):
    m = sum(xs) / len(xs)
    num = sum((a - m) * (b - m) for a, b in zip(xs, xs[1:]))
    den = sum((a - m) ** 2 for a in xs)
    return num / den if den else None


def a5(D, iv):
    R = Rows()
    if len(iv) < 10:
        return R
    for b in BLOCKS:
        lo, hi = block_boot(iv, b)
        R.add(f"A5_block{b}_lo", lo, "{:.2f}", f"T7 D, moving-block bootstrap, block {b}, 90% range, low")
        R.add(f"A5_block{b}_hi", hi, "{:.2f}", f"T7 D, moving-block bootstrap, block {b}, 90% range, high")
    # leads: growth and CPI inflation k quarters after each interval
    for k in LEADS:
        rows = []
        for x in iv:
            s3, e3 = L.month_add(x["start"], 3 * k), L.month_add(x["end"], 3 * k)
            qs = [q for q in D["gdp"] if s3 < T11.quarter_end(q) <= e3]
            g = (sum(D["gdp"][q] for q in qs) / len(qs)) if qs else None
            infl = []
            for m in L.months(L.month_add(s3, 1), e3):
                c, c0 = D["cpi"].get(m), D["cpi"].get(L.month_add(m, -12))
                if c and c0:
                    infl.append(100 * (c / c0 - 1))
            full = len(infl) == L.month_diff(s3, e3)
            rows.append(dict(x, g_lead=g, cpi_lead=(sum(infl) / len(infl)) if full and infl else None))
        for key, name in (("g_lead", "growth"), ("cpi_lead", "CPI inflation")):
            rp, rg, d, n = T11.rho_d(rows, "p", key)
            tag = f"A5_{'g' if key == 'g_lead' else 'cpi'}_lead{k}"
            R.add(tag + "_rho_p", rp, "{:.2f}", f"T7 with {name} led {k} quarter(s): rho with MAS's decisions")
            R.add(tag + "_rho", rg, "{:.2f}", f"T7 with {name} led {k} quarter(s): rho with {name}")
            R.add(tag + "_D", d, "{:.2f}", f"T7 with {name} led {k} quarter(s): D")
            R.add(tag + "_n", n, "{:d}", f"T7 with {name} led {k} quarter(s): intervals with data")
    # effective sample size, Bartlett: N (1 - r_a r_b) / (1 + r_a r_b), lag-1
    y, p, g = [x["y"] for x in iv], [float(x["p"]) for x in iv], [x["g"] for x in iv]
    ry, rp_, rg_ = lag1(y), lag1(p), lag1(g)
    R.add("A5_n", len(iv), "{:d}", "T7 intervals")
    R.add("A5_r1_y", ry, "{:.2f}", "lag-1 autocorrelation of the path y_i")
    R.add("A5_r1_p", rp_, "{:.2f}", "lag-1 autocorrelation of MAS's score p_i")
    R.add("A5_r1_g", rg_, "{:.2f}", "lag-1 autocorrelation of growth g_i")
    for nm, r in (("p", rp_), ("g", rg_)):
        ne = len(iv) * (1 - ry * r) / (1 + ry * r)
        R.add(f"A5_neff_{nm}", ne, "{:.0f}", f"effective sample size for rho(y, {nm}), Bartlett lag-1")
        R.add(f"A5_se_{nm}", 1 / math.sqrt(ne - 2), "{:.2f}", f"approximate standard error of rho(y, {nm}) at that size")
    return R


# ------------------------------------------------------------------ article forms
def context(D, W):
    """USD/JPY at the scored window's three-month endpoints, and the rank of
    Japan's real broad index (S1b) in its series."""
    R = Rows()
    first, last, start, end = W["scored"]
    u = D["usd"].get("JPY", {})
    for tag, ms in (("start", start), ("end", end)):
        if all(m in u for m in ms):
            v = math.exp(sum(math.log(u[m]) for m in ms) / len(ms))
            R.add(f"A6_usdjpy_{tag}", v, "{:.1f}", f"yen per US dollar, mean of the log over the scored window's {tag} months (S2)")
            R.add(f"A6_usdjpy_{tag}0", v, "{:.0f}", f"yen per US dollar, scored window {tag}, whole")
    re_ = D["reer"].get("JP", {})
    if re_:
        ps = sorted(re_)
        order = sorted(ps, key=lambda p: re_[p])
        R.add("A6_reer_jp_months", len(ps), "{:d}", f"months in Japan's real broad index, {ps[0]} to {ps[-1]}")
        R.add("A6_reer_jp_low_value", re_[order[0]], "{:.1f}", f"Japan's real broad index, lowest month ({order[0]})")
        R.add("A6_reer_jp_low_year", int(order[0][:4]), "{:d}", "year of that lowest month")
        if "2025-12" in re_:
            R.add("A6_reer_jp_2025_12_rank", order.index("2025-12") + 1, "{:d}",
                  "rank of December 2025 among all months, 1 = lowest")
        ann = {}
        for p in ps:
            ann.setdefault(p[:4], []).append(re_[p])
        full = {y: sum(v) / len(v) for y, v in ann.items() if len(v) == 12}
        ys = sorted(full, key=lambda y: full[y])
        R.add("A6_reer_jp_years", len(full), "{:d}", "full calendar years in Japan's real broad index")
        if "2025" in full:
            R.add("A6_reer_jp_2025_avg", full["2025"], "{:.1f}", "Japan's real broad index, 2025 average")
            R.add("A6_reer_jp_2025_rank", ys.index("2025") + 1, "{:d}", "rank of the 2025 average among full years, 1 = lowest")
        R.add("A6_reer_jp_lowest_year", int(ys[0]), "{:d}", "full year with the lowest average")
        R.add("A6_reer_jp_lowest_year_avg", full[ys[0]], "{:.1f}", "that year's average")
    return R


def article(D, sens, T):
    R = Rows()
    for m, tag in (("2021-01", "2021_01"), ("2025-12", "2025_12")):
        v = D["masfx"].get("JPY", {}).get(m)
        if v:
            R.add(f"art_sgd_per_100jpy_{tag}", 100 * v, "{:.3f}", f"MAS monthly average, S$ per 100 yen, {m}, three decimals")
            R.add(f"art_yen_per_budget_{tag}", round(1000 / v, -2), "{:,.0f}",
                  f"yen bought by S$1,000, {m}, to the nearest 100")
    S = {(r["variant"], r["key"]): r["value"] for r in sens}
    b1 = S.get(("single-month endpoints, scored", "single-month_scored_JPY_b"))
    if b1 not in (None, ""):
        R.add("art_single_month_JPY_pct0", 100 * (math.exp(float(b1)) - 1), "{:.0f}",
              "the Singapore dollar against the yen, single-month endpoints, per cent, whole")
    # Round 5: the specimen at the two-decimal rates the article prints.
    for m, tag in (("2021-01", "2021_01"), ("2025-12", "2025_12")):
        v = D["masfx"].get("JPY", {}).get(m)
        if v:
            r2 = round(100 * v, 2) / 100
            R.add(f"art_yen_per_budget_2dp_{tag}", 1000 / r2, "{:,.0f}", f"yen per S$1,000 at the two-decimal rate, {m}")
            R.add(f"art_yen_per_budget_2dp_k_{tag}", round(1000 / r2, -4 if m == "2025-12" else -3), "{:,.0f}",
                  f"the same, rounded as the article says it ('about'), {m}")
    for key, name in (("T2_P", "the yen's fall"), ("T2_S", "the Singapore dollar's rise")):
        v = T.get(key)
        if v not in (None, ""):
            R.add(f"art_{key}_pct10", round(10 * float(v)) * 10, "{:d}", f"T2 share, {name}, per cent to the nearest 10 ('about')")
    v = T.get("T1_long_JPY_R")
    if v not in (None, ""):
        R.add("art_long_JPY_R_pct0", abs(100 * float(v)), "{:.0f}", "long window: the yen split's remainder, per cent of the log change, negative")
    d = S.get(("finer slope coding", "D"))
    if d not in (None, ""):
        R.add("art_finer_D_3dp", float(d), "{:.3f}", "T7 with the finer slope coding: D, three decimals")
    return R


def main():
    a = L.args("sgd step 16: analyses added after outside review (not sealed, not scored)")
    P = L.paths(a.root)
    L.guard(P["raw"])
    out = P["out"]
    D = T11.load(out)
    D["masfx"] = D.get("masfx", {})
    W = T11.windows(D)
    iv = []
    for r in L.read_csv(os.path.join(out, "intervals.csv")):
        if r["y"] in ("", None) or r["g"] in ("", None):
            continue
        iv.append({"date": r["date"], "start": r["start"], "end": r["end"], "y": float(r["y"]),
                   "g": float(r["g"]), "p": int(r["p"])})
    sens = L.read_csv(os.path.join(out, "sensitivities.csv"))
    T = {r["key"]: r["value"] for r in L.read_csv(os.path.join(out, "tests.csv"))}
    for name, R in (("review_a1_usd.csv", a1(D, W)), ("review_a2_expair.csv", a2(D, W, P["raw"])),
                    ("review_a3_real.csv", a3(D, W, load_cpi_jp(P["raw"]))),
                    ("review_a4_rolling.csv", a4(D, out)), ("review_a5_t7.csv", a5(D, iv)),
                    ("review_a6_context.csv", context(D, W)),
                    ("review_article.csv", article(D, sens, T))):
        R.write(os.path.join(out, name))


if __name__ == "__main__":
    main()
